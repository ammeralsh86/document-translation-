import os
from docx import Document
from pptx import Presentation
import fitz  # PyMuPDF لقراءة ملفات PDF
from translator import DocumentTranslator

class DocumentHandler:
    def __init__(self, glossary_path="glossary.json"):
        """
        تهيئة معالج المستندات الشامل (Word, PowerPoint, PDF)
        """
        self.translator = DocumentTranslator(glossary_path=glossary_path)

    def translate_document(self, input_path, output_path, target_language="العربية"):
        """
        تحديد نوع الملف وتوجيهه للمعالج المناسب
        """
        if not os.path.exists(input_path):
            raise FileNotFoundError(f"الملف غير موجود في المسار المحدد: {input_path}")

        ext = os.path.splitext(input_path)[1].lower()

        if ext == ".docx":
            self._translate_docx(input_path, output_path, target_language)
        elif ext == ".pptx":
            self._translate_pptx(input_path, output_path, target_language)
        elif ext == ".pdf":
            self._translate_pdf_to_text(input_path, output_path, target_language)
        else:
            raise ValueError(f"امتداد الملف غير مدعوم حالياً: {ext}")

    def _translate_docx(self, input_path, output_path, target_language):
        print(f"[i] جاري معالجة وتعديل ملف Word: {input_path}")
        doc = Document(input_path)
        
        # ترجمة الفقرات
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                original_text = paragraph.text
                translated_text = self.translator.translate_text(original_text, target_language=target_language)
                if len(paragraph.runs) > 0:
                    paragraph.runs[0].text = translated_text
                    for run in paragraph.runs[1:]:
                        run.text = ""
                else:
                    paragraph.text = translated_text

        # ترجمة الجداول
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        cell.text = self.translator.translate_text(cell.text, target_language=target_language)

        doc.save(output_path)
        print(f"[✓] تم حفظ مستند Word المترجم في: {output_path}")

    def _translate_pptx(self, input_path, output_path, target_language):
        print(f"[i] جاري معالجة وترجمة شرائح الباوربوينت: {input_path}")
        prs = Presentation(input_path)

        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            original_text = paragraph.text
                            translated_text = self.translator.translate_text(original_text, target_language=target_language)
                            paragraph.text = translated_text

        prs.save(output_path)
        print(f"[✓] تم حفظ ملف الباوربوينت المترجم في: {output_path}")

    def _translate_pdf_to_text(self, input_path, output_path, target_language):
        """
        ملاحظة: ملفات PDF صعبة التعديل مباشرة بنفس التنسيق تماماً، 
        لذا نقوم باستخراج النصوص، ترجمتها، وحفظها كملف نصي أو مستخرج منظم.
        """
        print(f"[i] جاري قراءة واستخراج النصوص من ملف PDF: {input_path}")
        doc = fitz.open(input_path)
        translated_pdf_content = []

        for page_num, page in enumerate(doc):
            text = page.get_text()
            if text.strip():
                print(f"[i] جاري ترجمة الصفحة رقم {page_num + 1}...")
                translated_page = self.translator.translate_text(text, target_language=target_language)
                translated_pdf_content.append(f"--- صفحة {page_num + 1} ---\n{translated_page}\n")

        # حفظ النتيجة في ملف نصي أو مرن
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(translated_pdf_content))
            
        print(f"[✓] تم ترجمة محتوى الـ PDF بنجاح وحفظه في: {output_path}")

# للاختبار المباشر
if __name__ == "__main__":
    pass
